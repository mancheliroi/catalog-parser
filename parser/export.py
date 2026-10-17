"""Сохранение результатов в CSV и Excel."""

import csv
from dataclasses import asdict, fields
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from parser.extract import Product

HEADERS = {
    "title": "Название",
    "price": "Цена",
    "currency": "Валюта",
    "in_stock": "В наличии",
    "stock_count": "Остаток, шт.",
    "rating": "Рейтинг",
    "category": "Категория",
    "upc": "Артикул (UPC)",
    "url": "Ссылка",
    "description": "Описание",
}
COLUMNS = [f.name for f in fields(Product) if f.name in HEADERS]
COLUMNS.sort(key=list(HEADERS).index)


def _row(p: Product) -> dict:
    data = asdict(p)
    data["in_stock"] = "да" if p.in_stock else "нет"
    return {k: data[k] for k in COLUMNS}


def to_csv(products: list[Product], path: Path) -> None:
    # utf-8-sig — чтобы Excel сразу открыл кириллицу без кракозябр
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, delimiter=";")
        writer.writerow(HEADERS)
        writer.writerows(_row(p) for p in products)


def to_xlsx(products: list[Product], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"

    ws.append([HEADERS[c] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(vertical="center")

    for p in products:
        row = _row(p)
        ws.append([row[c] for c in COLUMNS])

    price_col = COLUMNS.index("price") + 1
    url_col = COLUMNS.index("url") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, price_col).number_format = "0.00"
        link = ws.cell(r, url_col)
        link.hyperlink = link.value
        link.font = Font(color="2563EB", underline="single")

    widths = {"title": 45, "description": 60, "url": 40}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save(products: list[Product], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".xlsx":
        to_xlsx(products, path)
    elif path.suffix.lower() == ".csv":
        to_csv(products, path)
    else:
        raise ValueError("Поддерживаются только .csv и .xlsx")

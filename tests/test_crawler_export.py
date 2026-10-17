import csv

import httpx
import pytest
from openpyxl import load_workbook

from parser.crawler import crawl, fetch
from parser.export import save

START = "https://books.toscrape.com/catalogue/page-1.html"


def client_for(site: dict[str, str]) -> httpx.AsyncClient:
    def handler(request: httpx.Request) -> httpx.Response:
        body = site.get(str(request.url))
        return httpx.Response(200, text=body) if body is not None else httpx.Response(404)

    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


@pytest.mark.asyncio
async def test_crawl_follows_pagination_dedupes_and_skips_broken(site):
    async with client_for(site) as client:
        products, failed = await crawl(START, client=client, delay=0)
    assert [p.title for p in products] == ["Книга А", "Книга Б"]
    assert failed == ["https://books.toscrape.com/catalogue/book-c/index.html"]


@pytest.mark.asyncio
async def test_crawl_respects_max_pages(site):
    async with client_for(site) as client:
        products, _ = await crawl(START, client=client, delay=0, max_pages=1)
    assert len(products) == 2


@pytest.mark.asyncio
async def test_fetch_retries_on_server_error():
    calls = {"n": 0}

    def handler(request):
        calls["n"] += 1
        return httpx.Response(503) if calls["n"] < 3 else httpx.Response(200, text="ok")

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        assert await fetch(client, "https://x.test/", backoff=0) == "ok"
    assert calls["n"] == 3


@pytest.mark.asyncio
async def test_fetch_does_not_retry_404():
    calls = {"n": 0}

    def handler(request):
        calls["n"] += 1
        return httpx.Response(404)

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        with pytest.raises(httpx.HTTPStatusError):
            await fetch(client, "https://x.test/", backoff=0)
    assert calls["n"] == 1


@pytest.mark.asyncio
async def test_export_xlsx_and_csv(site, tmp_path):
    async with client_for(site) as client:
        products, _ = await crawl(START, client=client, delay=0)

    xlsx = tmp_path / "out.xlsx"
    save(products, xlsx)
    ws = load_workbook(xlsx).active
    assert ws["A1"].value == "Название"
    assert ws["A2"].value == "Книга А" and ws["B2"].value == 51.77
    assert ws.max_row == 3

    path = tmp_path / "out.csv"
    save(products, path)
    rows = list(csv.reader(path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0][0] == "Название" and rows[2][0] == "Книга Б" and rows[2][3] == "нет"
